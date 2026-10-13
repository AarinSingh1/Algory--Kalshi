"""Build a workbook for pasting Bloomberg price tables by hand (when Excel/API downloads are blocked).

One sheet per security (paste the Table from a G chart or HP screen at A4) and a Checklist sheet that counts the
rows pasted into each. Run:  python bloomberg_excel/make_paste_template.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill

from make_templates import FONT, BOLD, TITLE, INPUT, FILL, HEAD, YEN_CORE, YEN_EXTRA, OIL_GENERICS, OIL_EQUITIES, MACRO

OUT = Path(__file__).resolve().parent / "templates" / "Bloomberg_Paste_Template.xlsx"
FIRST_PRIORITY = ["CL1 Comdty", "CO1 Comdty", "USO US Equity", "XLE US Equity", "OVX Index", "VIX Index",
                  "USGG3M Index", "USDJPY Curncy", "USDJPYV1M Curncy"]
MAX_ROW = 20000


def securities():
    """(group, ticker, description) for every simple series in bbg_pull.py, deduplicated, priority ones first."""
    groups = ([("Oil futures", t, d) for t, _, d in OIL_GENERICS]
              + [("Oil stocks/ETFs", f"{t} US Equity", d) for t, d in OIL_EQUITIES.items()]
              + [("Macro", t, d) for t, _, d in MACRO]
              + [("Yen", t, d) for t, _, d in YEN_CORE + YEN_EXTRA])
    seen, rows = set(), []
    for g, t, d in groups:
        if t not in seen:
            seen.add(t)
            rows.append((g, t, d))
    rank = {t: i for i, t in enumerate(FIRST_PRIORITY)}
    return sorted(rows, key=lambda r: rank.get(r[1], len(rank)))


def put(ws, ref, value, font=FONT, fill=None):
    c = ws[ref]
    c.value, c.font = value, font
    if fill:
        c.fill = fill


def main():
    wb = Workbook()
    ck = wb.active
    ck.title = "Checklist"
    put(ck, "A1", "Bloomberg price data: paste checklist", TITLE)
    steps = [
        "1. On the Terminal chart (G) screen, click the orange security box, type the ticker below, press Enter.",
        "2. Click Max (or 5Y), change '1 Min' to Daily, then click Table.",
        "3. Click in the table, Ctrl+A, Ctrl+C. Click 'Go' below, then click the yellow cell A4 on that sheet and Ctrl+V.",
        "4. Paste the table as it comes, headers included. Rows pasted updates by itself. Save often; upload to Claude.",
        "Start with Priority 1. Extra 1-3 sheets are for anything not on the list (type its ticker in B1 there).",
    ]
    for i, s in enumerate(steps, 3):
        put(ck, f"A{i}", s)
    hdr = 9
    for i, h in enumerate(["#", "Priority", "Group", "Ticker", "Description", "Sheet", "Rows pasted", "Status"]):
        put(ck, f"{'ABCDEFGH'[i]}{hdr}", h, BOLD, HEAD)

    rows = securities() + [("Extra", "", f"Extra {i}") for i in (1, 2, 3)]
    for n, (g, t, d) in enumerate(rows, 1):
        name = t if t else d
        ws = wb.create_sheet(name)
        put(ws, "A1", "Ticker:", BOLD)
        put(ws, "B1", t, INPUT if not t else BOLD, FILL if not t else None)
        put(ws, "C1", d)
        put(ws, "A2", "Click the yellow cell A4 and press Ctrl+V to paste the Bloomberg table (headers included).")
        put(ws, "A3", "Paste below", BOLD)
        put(ws, "A4", None, INPUT, FILL)
        for col, w in zip("ABCDEFGH", (14, 16, 14, 14, 14, 14, 14, 14)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A5"

        r = hdr + n
        ref = f"'{name}'"
        put(ck, f"A{r}", n)
        put(ck, f"B{r}", 1 if t in FIRST_PRIORITY else 2)
        put(ck, f"C{r}", g)
        put(ck, f"D{r}", f'=IF({ref}!B1="","",{ref}!B1)' if not t else t)
        put(ck, f"E{r}", d)
        put(ck, f"F{r}", f'=HYPERLINK("#{ref}!A4","Go")', Font(name="Arial", size=10, color="0563C1", underline="single"))
        put(ck, f"G{r}", f"=COUNTA({ref}!A5:A{MAX_ROW})")
        put(ck, f"H{r}", f'=IF(G{r}>0,"Done","To do")')
    last = hdr + len(rows)
    ck.conditional_formatting.add(f"H{hdr + 1}:H{last}", CellIsRule(
        operator="equal", formula=['"Done"'], fill=PatternFill("solid", fgColor="C6EFCE")))
    put(ck, f"A{last + 2}", "Done so far:", BOLD)
    put(ck, f"C{last + 2}", f'=COUNTIF(H{hdr + 1}:H{last},"Done")&" of "&COUNTA(A{hdr + 1}:A{last})')
    for col, w in zip("ABCDEFGH", (5, 9, 16, 22, 45, 7, 12, 9)):
        ck.column_dimensions[col].width = w
    ck.freeze_panes = f"A{hdr + 1}"
    wb.save(OUT)
    print(f"wrote {OUT.name}: {len(rows)} paste sheets")


if __name__ == "__main__":
    main()
